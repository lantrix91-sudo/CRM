from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import uuid

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.customers.models import Client
from apps.leads.models import TelegramNotice
from apps.services.models import Service
from .models import Order, WorkerTransfer
from .services import cancel_order


class RegressionTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username="manager", role="manager")
        self.worker = User.objects.create_user(username="worker", role="worker", percentage=40, is_available=True)
        self.service = Service.objects.create(name="Repair")
        self.worker.services.add(self.service)
        self.order = Order.objects.create(
            title="Repair", client=Client.objects.create(name="Client", phone="+77001234567"),
            service=self.service, amount=1000,
        )
        self.client.force_login(self.manager)

    def save_order(self, **changes):
        data = dict(action="save", title="Updated", client=self.order.client_id,
                    service=self.service.pk, employee=self.worker.pk, amount="1000")
        data.update(changes)
        return self.client.post(f"/orders/{self.order.pk}/", data)

    def paid_order(self):
        self.order.employee = self.worker
        self.order.status = "paid"
        self.order.worker_percentage = 40
        self.order.paid_at = timezone.now()
        self.order.save()

    def export(self, **params):
        return self.client.get(f"/settlements/{self.worker.pk}/export/", params)

    def test_card_assignment_creates_notice_and_event(self):
        self.assertEqual(self.save_order().status_code, 302)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "assigned")
        self.assertEqual(self.order.employee, self.worker)
        self.assertEqual(TelegramNotice.objects.filter(order=self.order, active=True).count(), 1)
        self.assertEqual(self.order.events.count(), 1)

    def test_invalid_assignment_rolls_back_other_edits(self):
        for condition in ("unskilled", "unavailable", "full"):
            with self.subTest(condition=condition):
                self.worker.services.set([] if condition == "unskilled" else [self.service])
                self.worker.is_available = condition != "unavailable"
                self.worker.max_active_leads = 1 if condition == "full" else 5
                if condition == "full":
                    Order.objects.create(title="Busy", client=self.order.client, service=self.service,
                                         employee=self.worker, status="in_progress")
                self.worker.save()
                self.assertEqual(self.save_order().status_code, 200)
                self.order.refresh_from_db()
                self.assertIsNone(self.order.employee_id)
                self.assertEqual(self.order.title, "Repair")
                self.assertEqual(self.order.status, "new")
                self.assertFalse(TelegramNotice.objects.filter(order=self.order).exists())

    def test_assigned_card_cannot_replace_worker_or_service(self):
        self.save_order()
        other = User.objects.create_user(username="other", role="worker")
        other_service = Service.objects.create(name="Other")
        self.assertEqual(self.save_order(employee=other.pk, service=other_service.pk).status_code, 302)
        self.order.refresh_from_db()
        self.assertEqual(self.order.employee, self.worker)
        self.assertEqual(self.order.service, self.service)
        self.assertEqual(TelegramNotice.objects.filter(order=self.order, active=True).count(), 1)

    def test_cancel_preserves_full_reason_and_event(self):
        reason = "x" * 300
        cancel_order(self.order.pk, self.manager, reason)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "cancelled")
        self.assertEqual(self.order.cancellation_reason, reason)
        self.assertTrue(self.order.events.get().description.endswith(reason))

    def test_export_comments_are_text_and_times_are_local(self):
        self.paid_order()
        self.order.work_comment = "=1+1"
        self.order.save()
        WorkerTransfer.objects.create(worker=self.worker, received_by=self.manager,
                                     amount=1, request_id=uuid.uuid4(), comment="=2+2")
        response = self.export()
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2, 12).value, "=1+1")
        self.assertEqual(sheet.cell(3, 12).value, "=2+2")
        self.assertEqual(sheet.cell(2, 12).data_type, "s")
        self.assertEqual(sheet.cell(3, 12).data_type, "s")
        self.assertEqual(sheet.cell(2, 1).value, timezone.localtime(self.order.paid_at).strftime("%Y-%m-%d %H:%M"))

    def test_invalid_export_parameters_return_400(self):
        for params in ({"date_from": "invalid"}, {"date_from": "2026-02-02", "date_to": "2026-01-01"}, {"date_to": "9999-12-31"}):
            with self.subTest(params=params):
                self.assertEqual(self.export(**params).status_code, 400)
        for worker in ("", "abc", "9" * 100):
            self.assertEqual(self.client.get("/settlements/export/", {"worker": worker}).status_code, 400)

    def test_missing_percentage_returns_400_including_opening_balance(self):
        self.paid_order()
        self.order.worker_percentage = None
        self.order.paid_at = timezone.now() - timedelta(days=2)
        self.order.save()
        self.assertEqual(self.export().status_code, 400)
        self.assertEqual(self.export(date_from=timezone.localdate().isoformat()).status_code, 400)


    def test_free_repeat_has_no_payment_controls_and_rejects_payment(self):
        self.paid_order()
        repeated = Order.objects.create(
            title="Repeat", client=self.order.client, service=self.service,
            employee=self.worker, repeat_of=self.order, status="completed", amount=1000,
        )
        url = f"/orders/{repeated.pk}/"
        for user, action in ((self.worker, "received_payment"), (self.manager, "pay")):
            with self.subTest(role=user.role):
                self.client.force_login(user)
                response = self.client.get(url)
                self.assertContains(response, "Оплата не требуется.")
                self.assertNotContains(response, 'value="received_payment"')
                self.assertNotContains(response, 'value="pay"')
                self.assertNotContains(response, "Оплата от клиента")
                response = self.client.post(url, {
                    "action": action, "received_amount": "1000", "received_method": "cash",
                })
                self.assertEqual(response.status_code, 200)
                repeated.refresh_from_db()
                self.assertEqual(repeated.status, "completed")
                self.assertIsNone(repeated.received_at)
                self.assertIsNone(repeated.paid_at)


    def test_profile_history_filters_orders_and_keeps_worker_scope(self):
        self.paid_order()
        older = Order.objects.create(title="Old", client=self.order.client, service=self.service,
            employee=self.worker, status="paid", amount=200, worker_percentage=40,
            paid_at=timezone.now() - timedelta(days=2))
        other = User.objects.create_user(username="other_history", role="worker")
        private = Order.objects.create(title="Private", client=self.order.client, service=self.service,
            employee=other, status="paid", amount=300, worker_percentage=40, paid_at=timezone.now())
        self.client.force_login(self.worker)
        today = self.client.get("/profile/?history=today")
        self.assertEqual(today.status_code, 200)
        self.assertEqual([o.pk for o in today.context["orders"]], [self.order.pk])
        history = self.client.get(f"/profile/?history=earnings&worker={other.pk}")
        self.assertEqual({o.pk for o in history.context["orders"]}, {self.order.pk, older.pk})
        self.assertNotIn(private.pk, [o.pk for o in history.context["orders"]])
        self.assertContains(history, f'/orders/{self.order.pk}/#order-history')
        Order.objects.bulk_create([Order(title=f"Extra {i}", client=self.order.client,
            service=self.service, employee=self.worker, status="paid", amount=100,
            worker_percentage=40, paid_at=timezone.now()) for i in range(21)])
        page = self.client.get("/profile/?history=earnings&page=2")
        self.assertEqual(page.context["page_obj"].number, 2)
        self.assertEqual(len(page.context["orders"]), 3)


    def test_manager_stages_exclude_completed_free_repeats(self):
        self.order.status = "completed"
        self.order.save()
        for index in range(2):
            Order.objects.create(title=f"Repeat {index}", client=self.order.client,
                service=self.service, employee=self.worker, repeat_of=self.order,
                status="completed")
        response = self.client.get("/manager/")
        stages = {stage["label"]: stage["total"] for stage in response.context["stages"]}
        self.assertEqual(stages[Order.Status.COMPLETED.label], 1)
        self.assertNotIn("Выполнен · бесплатно (повторный ремонт)", stages)
        self.assertEqual(sum(stages.values()), 1)
        self.order.status = "paid"
        self.order.save()
        response = self.client.get("/manager/")
        stages = {stage["label"]: stage["total"] for stage in response.context["stages"]}
        self.assertNotIn(Order.Status.COMPLETED.label, stages)
        self.assertNotIn("Выполнен · бесплатно (повторный ремонт)", stages)


    def test_zero_completion_is_free_and_cannot_be_paid(self):
        from .services import complete_order_with_payment
        self.order.employee = self.worker
        self.order.status = "in_progress"
        self.order.save()
        complete_order_with_payment(self.order.pk, Decimal("0"), Decimal("0"), "No charge", actor=self.worker)
        self.order.refresh_from_db()
        self.assertTrue(self.order.is_free)
        self.assertEqual(self.order.amount, 0)
        self.assertEqual(self.order.status, "completed")
        for user, action in ((self.worker, "received_payment"), (self.manager, "pay")):
            self.client.force_login(user)
            url = f"/orders/{self.order.pk}/"
            page = self.client.get(url)
            self.assertContains(page, "Оплата не требуется.")
            self.assertNotContains(page, 'value="received_payment"')
            self.assertNotContains(page, 'value="pay"')
            self.client.post(url, {"action": action, "received_amount": "100", "received_method": "cash"})
            self.order.refresh_from_db()
            self.assertEqual(self.order.status, "completed")
            self.assertIsNone(self.order.received_at)
        self.client.force_login(self.worker)
        self.assertNotContains(self.client.get("/my-orders/"), f'/orders/{self.order.pk}/')
        self.assertContains(self.client.get("/my-orders/history/"), f'/orders/{self.order.pk}/')
