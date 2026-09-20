from apps.orders.test_helpers import create_order_with_rate
from decimal import Decimal
import uuid
from datetime import timedelta
from openpyxl import load_workbook
from django.test import TestCase
from django.core.exceptions import ValidationError, PermissionDenied
from apps.accounts.models import User, WorkerServiceRate
from apps.customers.models import Client
from apps.services.models import Service
from .models import Order, WorkerTransfer, SettlementShift
from .settlements import settle, totals
from .reporting import export_rows

class SettlementTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username="manager", role="manager")
        self.worker = User.objects.create_user(username="worker", role="worker")
        self.order = create_order_with_rate(title="Repair", client=Client.objects.create(name="Client", phone="123"),
            service=Service.objects.create(name="Repair"), employee=self.worker, status="paid", amount=20000, expenses=2000, worker_percentage=40)

    def test_partial_transfer_close_and_carry(self):
        summary = totals(self.worker)
        self.assertEqual(summary["current_shift_gross"], Decimal("20000"))
        self.assertEqual(summary["current_shift_worker_share"], Decimal("7200"))
        self.assertEqual(summary["current_shift_company_share"], Decimal("10800"))
        self.assertEqual(summary["carried_debt"], Decimal("0"))
        self.assertEqual(summary["current_shift_debt"], Decimal("10800"))
        self.assertEqual(summary["total_balance_due"], Decimal("10800"))
        token = uuid.uuid4()
        settle(self.worker.pk, self.manager, "transfer", Decimal("8000"), request_id=token)
        settle(self.worker.pk, self.manager, "transfer", Decimal("8000"), request_id=token)
        self.assertEqual(WorkerTransfer.objects.count(), 1)
        settle(self.worker.pk, self.manager, "close")
        summary = totals(self.worker)
        self.assertEqual(summary["current_shift_company_share"], Decimal("0"))
        self.assertEqual(summary["current_shift_worker_share"], Decimal("0"))
        self.assertEqual(summary["carried_debt"], Decimal("2800"))
        self.assertEqual(summary["current_shift_debt"], Decimal("0"))
        self.assertEqual(summary["total_balance_due"], Decimal("2800"))
        with self.assertRaises(ValidationError):
            settle(self.worker.pk, self.manager, "close")
        settle(self.worker.pk, self.manager, "transfer", Decimal("2800"), request_id=uuid.uuid4())
        summary = totals(self.worker)
        self.assertEqual(summary["carried_debt"], Decimal("0"))
        self.assertEqual(summary["total_balance_due"], Decimal("0"))
        self.assertEqual(SettlementShift.objects.get().balance, 2800)
        self.order.refresh_from_db()
        self.assertEqual(self.order.amount, 20000)

    def test_new_order_is_current_and_carried_debt_is_not_double_counted(self):
        settle(self.worker.pk, self.manager, "transfer", Decimal("4300"), request_id=uuid.uuid4())
        settle(self.worker.pk, self.manager, "close")

        new_order = create_order_with_rate(
            title="New repair", client=self.order.client, service=self.order.service,
            employee=self.worker, status="paid", amount=15000, worker_percentage=40,
        )
        summary = totals(self.worker)
        self.assertEqual(summary["carried_debt"], Decimal("6500"))
        self.assertEqual(summary["current_shift_gross"], Decimal("15000"))
        self.assertEqual(summary["current_shift_company_share"], Decimal("9000"))
        self.assertEqual(summary["current_shift_debt"], Decimal("9000"))
        self.assertEqual(summary["total_balance_due"], Decimal("15500"))
        self.assertEqual(summary["total_transferred_all_time"], Decimal("4300"))

        settle(self.worker.pk, self.manager, "transfer", Decimal("6500"), request_id=uuid.uuid4())
        summary = totals(self.worker)
        self.assertEqual(summary["carried_debt"], Decimal("0"))
        self.assertEqual(summary["current_shift_debt"], Decimal("9000"))
        self.assertEqual(summary["total_balance_due"], Decimal("9000"))
        self.assertEqual(new_order.settlement_shift_id, None)

    def test_closing_fully_paid_shift_starts_without_carried_debt(self):
        settle(self.worker.pk, self.manager, "transfer", Decimal("10800"), request_id=uuid.uuid4())
        settle(self.worker.pk, self.manager, "close")
        summary = totals(self.worker)
        self.assertEqual(summary["carried_debt"], Decimal("0"))
        self.assertEqual(summary["total_balance_due"], Decimal("0"))
        self.assertEqual(SettlementShift.objects.get().balance, Decimal("0"))

    def test_closed_shift_detail_and_excel_export_are_restricted_and_historical(self):
        settle(self.worker.pk, self.manager, "close")
        shift = SettlementShift.objects.get()
        self.client.force_login(self.worker)

        detail = self.client.get(f"/settlements/shifts/{shift.pk}/")
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, f"Закрытая смена № {shift.pk}")
        self.assertContains(detail, "Доля компании: 10800,00 KZT")
        WorkerServiceRate.objects.create(worker=self.worker,
            service=Service.objects.create(name="Other rate"), worker_percentage=80)
        self.assertContains(self.client.get(f"/settlements/shifts/{shift.pk}/"), "Процент мастера: 40,00")

        export = self.client.get("/settlements/export/")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(export["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(filename=__import__("io").BytesIO(export.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Дата")
        self.assertEqual(rows[-1][0], "Остаток на конец периода")
        self.assertEqual(rows[-4][8], Decimal("10800.00"))

        other = User.objects.create_user(username="other_worker", role="worker")
        self.assertEqual(self.client.get(f"/settlements/shifts/{other.pk}/").status_code, 404)

    def test_export_uses_opening_balance_and_chronological_partial_transfers(self):
        settle(self.worker.pk, self.manager, "transfer", Decimal("4300"), request_id=uuid.uuid4())
        settle(self.worker.pk, self.manager, "close")
        shift = SettlementShift.objects.get()
        from django.utils import timezone
        now = timezone.now()
        SettlementShift.objects.filter(pk=shift.pk).update(closed_at=now - timedelta(days=2))
        Order.objects.filter(pk=self.order.pk).update(paid_at=now - timedelta(days=2))
        WorkerTransfer.objects.filter(worker=self.worker, amount=Decimal("4300")).update(
            created_at=now - timedelta(days=2)
        )
        client = self.order.client
        service = Service.objects.create(name="Company-only service")
        order_72 = create_order_with_rate(
            title="Order 72", client=client, service=service, employee=self.worker,
            status="paid", amount=10000, worker_percentage=0, paid_at=now,
        )
        order_73 = create_order_with_rate(
            title="Order 73", client=client, service=service, employee=self.worker,
            status="paid", amount=7500, worker_percentage=0, paid_at=now + timedelta(minutes=2),
        )
        order_74 = create_order_with_rate(
            title="Order 74", client=client, service=Service.objects.create(name="Shared service"), employee=self.worker,
            status="paid", amount=5000, expenses=2000, worker_percentage=50,
            paid_at=now + timedelta(minutes=4),
        )
        transfer_1 = WorkerTransfer.objects.create(
            worker=self.worker, received_by=self.manager, amount=16500,
            request_id=uuid.uuid4(), created_at=now + timezone.timedelta(minutes=1),
        )
        transfer_2 = WorkerTransfer.objects.create(
            worker=self.worker, received_by=self.manager, amount=7500,
            request_id=uuid.uuid4(),
        )
        transfer_3 = WorkerTransfer.objects.create(
            worker=self.worker, received_by=self.manager, amount=200,
            request_id=uuid.uuid4(),
        )
        WorkerTransfer.objects.filter(pk=transfer_1.pk).update(created_at=now + timedelta(minutes=1))
        WorkerTransfer.objects.filter(pk=transfer_2.pk).update(created_at=now + timedelta(minutes=3))
        WorkerTransfer.objects.filter(pk=transfer_3.pk).update(created_at=now + timedelta(minutes=5))
        rows = export_rows(self.worker, now.date())
        self.assertEqual(rows[0]["operation"], "Входящий остаток")
        self.assertEqual(rows[0]["remaining_balance"], Decimal("6500"))
        self.assertEqual([row["operation"] for row in rows[1:]], [
            "Оплата заказа", "Перевод от мастера", "Оплата заказа",
            "Перевод от мастера", "Оплата заказа", "Перевод от мастера",
        ])
        self.assertEqual(rows[-1]["remaining_balance"], Decimal("1300"))
        self.assertEqual(sum(row["company_amount"] for row in rows), Decimal("19000"))
        self.assertEqual(sum(row["transfer_amount"] for row in rows), Decimal("24200"))

    def test_export_keeps_order_and_transfer_comments_in_their_own_rows(self):
        from django.utils import timezone

        self.order.work_comment = "Заменил сливной насос, почистил фильтр"
        self.order.paid_at = timezone.now()
        self.order.save(update_fields=("work_comment", "paid_at"))
        transfer = WorkerTransfer.objects.create(
            worker=self.worker,
            received_by=self.manager,
            amount=Decimal("100"),
            comment="Перевод за наличные расходы",
            request_id=uuid.uuid4(),
        )
        rows = export_rows(self.worker)
        order_row = next(row for row in rows if row["order_number"] == self.order.pk)
        transfer_row = next(row for row in rows if row["transfer_amount"] == transfer.amount)
        self.assertEqual(order_row["comment"], self.order.work_comment)
        self.assertEqual(transfer_row["comment"], transfer.comment)
        self.assertNotEqual(order_row["comment"], transfer_row["comment"])

        self.order.work_comment = ""
        self.order.save(update_fields=("work_comment",))
        transfer.comment = ""
        transfer.save(update_fields=("comment",))
        rows = export_rows(self.worker)
        order_row = next(row for row in rows if row["order_number"] == self.order.pk)
        transfer_row = next(row for row in rows if row["transfer_amount"] == transfer.amount)
        self.assertEqual(order_row["comment"], "")
        self.assertEqual(transfer_row["comment"], "")

    def test_repeat_repairs_are_excluded_from_settlement_and_export_financials(self):
        repeat = Order.objects.create(
            title="Free repeat",
            client=self.order.client,
            service=self.order.service,
            employee=self.worker,
            repeat_of=self.order,
            status="completed",
        )
        before = totals(self.worker)
        self.assertEqual(before["current_shift_gross"], Decimal("20000"))
        self.assertEqual(before["current_shift_company_share"], Decimal("10800"))
        rows = export_rows(self.worker)
        self.assertNotIn(repeat.pk, [row["order_number"] for row in rows])

    def test_validation_and_access(self):
        with self.assertRaises(PermissionDenied):
            settle(self.worker.pk, self.worker, "close")
        with self.assertRaises(ValidationError):
            settle(self.worker.pk, self.manager, "transfer", Decimal("20000"), request_id=uuid.uuid4())
        self.client.force_login(self.worker)
        self.assertEqual(self.client.get("/settlements/").status_code, 200)
        self.assertEqual(self.client.post(f"/settlements/{self.worker.pk}/", {"action":"close"}).status_code, 403)
        other = User.objects.create_user(username="other", role="worker")
        self.assertEqual(self.client.get(f"/settlements/{other.pk}/").status_code, 403)
        self.client.force_login(self.manager)
        self.assertEqual(self.client.post(f"/settlements/{self.worker.pk}/", {"action":"transfer", "amount":"8000", "request_id":str(uuid.uuid4())}).status_code, 302)
        self.assertEqual(totals(self.worker)["balance"], 2800)
        self.worker.service_rates.filter(service=self.order.service).delete()
        with self.assertRaises(ValidationError):
            settle(self.worker.pk, self.manager, "close")

    def test_settlement_page_limits_history_but_export_keeps_older_records(self):
        for index in range(6):
            WorkerTransfer.objects.create(
                worker=self.worker,
                received_by=self.manager,
                amount=Decimal(index + 1),
                request_id=uuid.uuid4(),
                comment=f"transfer-{index + 1}",
            )
            SettlementShift.objects.create(
                worker=self.worker,
                closed_by=self.manager,
                manager_amount=Decimal("100"),
                worker_amount=Decimal("100"),
                balance=Decimal(index),
            )

        self.client.force_login(self.worker)
        response = self.client.get("/settlements/")
        self.assertEqual(response.status_code, 200)
        recent_transfers = list(response.context["recent_transfers"])
        recent_shifts = list(response.context["recent_shifts"])
        self.assertEqual(len(recent_transfers), 5)
        self.assertEqual(len(recent_shifts), 5)
        self.assertEqual(recent_transfers, list(self.worker.settlement_transfers.order_by("-created_at", "-pk")[:5]))
        self.assertEqual(recent_shifts, list(self.worker.closed_shifts.order_by("-closed_at", "-pk")[:5]))
        self.assertNotContains(response, "transfer-1")
        self.assertContains(response, "transfer-6")
        oldest_shift = self.worker.closed_shifts.order_by("closed_at", "pk").first()
        self.assertNotContains(response, f"Смена № {oldest_shift.pk}")

        self.assertEqual(WorkerTransfer.objects.filter(worker=self.worker).count(), 6)
        self.assertEqual(SettlementShift.objects.filter(worker=self.worker).count(), 6)
        export = self.client.get("/settlements/export/")
        workbook = load_workbook(filename=__import__("io").BytesIO(export.content), read_only=True)
        values = [value for row in workbook.active.iter_rows(values_only=True) for value in row]
        self.assertIn("transfer-1", values)
