from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Count, Sum, Q
from django.shortcuts import get_object_or_404, redirect, render
from decimal import Decimal

from .access import allowed, role_of, roles_required
from .forms import ReceivedPaymentForm, OrderCompletionForm, ClientForm, LeadForm, OrderForm, EmployeeForm, ServiceForm
from .models import User, WorkerServiceRate
from apps.leads.models import Lead, TelegramNotice
from apps.orders.models import Order
from apps.services.models import Service
from apps.customers.models import City
from apps.orders.services import convert_lead, transition_order, complete_order_with_payment, record_event

@login_required
def home(request):
    role = role_of(request.user)
    if role not in ("manager", "operator", "worker", "curator"):
        return render(request, "accounts/unassigned.html", status=403)
    return redirect({"manager": "manager-home", "operator": "operator-home", "worker": "worker-home", "curator": "profile"}[role])

@roles_required("manager")
def manager_home(request):
    stages = [
        {"label": Order.Status(row["status"]).label, "total": row["total"]}
        for row in Order.objects.exclude(Q(status="completed") & (Q(repeat_of__isnull=False) | Q(is_free=True)))
        .values("status").annotate(total=Count("pk")).order_by("status")
    ]
    return render(request, "accounts/manager.html", {
        "stages": stages,
        "workers": User.objects.filter(role="worker").annotate(active_orders=Count("orders", filter=Q(orders__status__in=("new", "assigned", "in_progress")), distinct=True), active_leads=Count("leads", filter=Q(leads__status__in=("new", "assigned", "in_progress")), distinct=True)).order_by("username"),
        "paid": Order.objects.filter(status="paid").aggregate(total=Sum("amount"))["total"] or 0,
        "leads_count": Lead.objects.count(),
        "services": Service.objects.order_by("name"),
    })

@roles_required("manager", "operator")
def operator_home(request):
    return render(request, "accounts/operator.html", {
        "leads": Lead.objects.exclude(status__in=("converted", "lost")).select_related("client", "service").order_by("-created_at")[:100],
        "orders": Order.objects.select_related("client", "employee").order_by("-created_at")[:100],
    })

@roles_required("worker")
def worker_home(request):
    return render(request, "accounts/worker.html", {
        "assigned_leads": Lead.objects.filter(employee=request.user, order__isnull=True, status__in=("new", "assigned", "in_progress")).select_related("client", "service").order_by("-created_at")[:100],
        "orders": Order.objects.filter(employee=request.user).exclude(status__in=("paid", "cancelled")).exclude(Q(status="completed") & (Q(is_free=True) | Q(repeat_of__isnull=False))).select_related("client", "service").order_by("-created_at")[:100],
    })

@roles_required("manager", "operator", "worker", "curator")
def order_detail(request, pk):
    with transaction.atomic():
        queryset = Order.objects.all()
        if role_of(request.user) == "worker":
            queryset = queryset.filter(employee=request.user)
        if role_of(request.user) == "curator":
            queryset = queryset.filter(employee__curator=request.user)
            if request.method != "GET":
                raise PermissionDenied
        order = get_object_or_404(queryset.select_for_update(), pk=pk)
        form = None
        payment_form = ReceivedPaymentForm(request.POST if request.method == "POST" and request.POST.get("action") == "received_payment" else None)
        preview_key = f"order_completion_preview:{pk}"
        preview = request.session.get(preview_key)
        completion_form = OrderCompletionForm(
            request.POST if request.method == "POST" and request.POST.get("action") in (
                "complete_with_payment", "preview_completion",
            ) else None,
            initial=preview if preview and request.method == "GET" else None,
        )
        if allowed(request.user, "manager", "operator") and order.status != "cancelled":
            form = OrderForm(request.POST if request.method == "POST" and request.POST.get("action") == "save" else None, instance=order)
        if request.method == "POST":
            action = request.POST.get("action")
            try:
                if action == "repeat_repair":
                    from apps.orders.services import repeat_repair
                    repeated = repeat_repair(pk, request.user)
                    messages.success(request, f"Повторный ремонт № {repeated.pk} назначен мастеру. Уведомление поставлено в очередь Telegram.")
                    return redirect("kanban")
                elif action == "cancel":
                    from apps.orders.services import cancel_order
                    cancel_order(pk, request.user, request.POST.get("reason", ""))
                    if request.POST.get("return_to") == "kanban":
                        return redirect("kanban")
                    return redirect("order-detail", pk=pk)
                elif action == "save":
                    if form is None:
                        raise PermissionDenied
                    previous_employee = order.employee_id
                    if form.is_valid():
                        # A savepoint rolls back field edits if assignment validation fails.
                        with transaction.atomic():
                            order = form.save(commit=False)
                            employee_id = order.employee_id
                            order.employee_id = previous_employee
                            order.save()
                            if employee_id != previous_employee:
                                from apps.orders.services import assign_order
                                order = assign_order(pk, employee_id, request.user)
                        return redirect("order-detail", pk=pk)
                elif action == "received_payment":
                    if role_of(request.user) != "worker" or order.employee_id != request.user.pk:
                        raise PermissionDenied
                    if order.repeat_of_id or order.is_free:
                        raise ValidationError("Повторный ремонт выполняется бесплатно и не требует оплаты.")
                    if order.status != "completed":
                        raise ValidationError("Сообщить об оплате можно после завершения работы.")
                    if payment_form.is_valid():
                        from django.utils import timezone
                        order.received_amount = payment_form.cleaned_data["received_amount"]
                        order.received_method = payment_form.cleaned_data["received_method"]
                        order.received_at = timezone.now()
                        order.amount = order.received_amount
                        order.save(update_fields=("received_amount", "received_method", "received_at", "amount"))
                        transition_order(pk, "pay", actor=request.user)
                        record_event(order, request.user, f"Сообщил о получении оплаты: {order.received_amount} KZT, {order.get_received_method_display()}")
                        messages.success(request, "Оплата сохранена. Заказ оплачен.")
                        return redirect("order-detail", pk=pk)
                elif action == "preview_completion":
                    if role_of(request.user) != "worker" or order.employee_id != request.user.pk:
                        raise PermissionDenied
                    if order.status != Order.Status.IN_PROGRESS or order.repeat_of_id:
                        raise ValidationError("Предпросмотр завершения больше недоступен.")
                    if completion_form.is_valid():
                        from apps.orders.services import payment_split, worker_service_percentage
                        amount = completion_form.cleaned_data["amount"]
                        expenses = completion_form.cleaned_data["expenses"]
                        percentage = worker_service_percentage(order.employee, order.service)
                        net, worker_amount, manager_amount = payment_split(
                            order, amount=amount, expenses=expenses,
                            worker_percentage=percentage,
                        )
                        request.session[preview_key] = {
                            "amount": str(amount), "expenses": str(expenses),
                            "comment": completion_form.cleaned_data["comment"],
                            "status": order.status, "employee_id": order.employee_id,
                            "worker_percentage": str(percentage) if percentage is not None else None,
                            "net": str(net), "worker_amount": str(worker_amount) if worker_amount is not None else None,
                            "company_amount": str(manager_amount) if manager_amount is not None else None,
                        }
                        request.session.modified = True
                        preview = request.session[preview_key]
                    else:
                        raise ValidationError("Проверьте сумму, расходы и комментарий.")
                elif action == "completion_back":
                    request.session.pop(preview_key, None)
                    request.session.modified = True
                    completion_form = OrderCompletionForm(initial=preview or None)
                    preview = None
                elif action == "confirm_completion":
                    if role_of(request.user) != "worker" or order.employee_id != request.user.pk:
                        raise PermissionDenied
                    if not preview or order.status != Order.Status.IN_PROGRESS or order.employee_id != preview.get("employee_id"):
                        request.session.pop(preview_key, None)
                        request.session.modified = True
                        raise ValidationError("Предпросмотр устарел. Заполните данные заново.")
                    from apps.orders.services import worker_service_percentage
                    current_rate = worker_service_percentage(order.employee, order.service)
                    current_percentage = str(current_rate) if current_rate is not None else None
                    if current_percentage != preview.get("worker_percentage"):
                        request.session.pop(preview_key, None)
                        request.session.modified = True
                        raise ValidationError("Процент мастера изменился. Создайте новый предпросмотр.")
                    complete_order_with_payment(
                        pk, Decimal(preview["amount"]), Decimal(preview["expenses"]),
                        preview["comment"], actor=request.user,
                    )
                    request.session.pop(preview_key, None)
                    request.session.modified = True
                    messages.success(request, "Заказ завершён.")
                    return redirect("order-detail", pk=pk)
                elif action == "complete_with_payment":
                    if role_of(request.user) != "worker" or order.employee_id != request.user.pk:
                        raise PermissionDenied
                    if order.repeat_of_id:
                        complete_order_with_payment(
                            pk, 0, 0, "Повторка выполнена бесплатно", actor=request.user,
                        )
                    else:
                        raise ValidationError("Сначала просмотрите и подтвердите завершение заказа.")
                    messages.success(request, "Заказ завершён.")
                    return redirect("order-detail", pk=pk)
                elif action == "pay":
                    if not allowed(request.user, "manager"):
                        raise PermissionDenied
                    transition_order(pk, "pay", actor=request.user)
                    return redirect("order-detail", pk=pk)
                elif action == "start":
                    if role_of(request.user) != "worker" or order.employee_id != request.user.pk:
                        raise PermissionDenied
                    transition_order(pk, action, actor=request.user)
                    return redirect("order-detail", pk=pk)
                elif action == "reject":
                    if role_of(request.user) != "worker" or order.status != "assigned":
                        raise PermissionDenied
                    record_event(order, request.user, "Мастер отклонил заказ")
                    order.employee = None
                    order.status = "new"
                    order.save(update_fields=("employee", "status"))
                    TelegramNotice.objects.filter(order=order, active=True).update(active=False)
                    return redirect("worker-home")
                else:
                    raise PermissionDenied
            except ValidationError as error:
                messages.error(request, "; ".join(error.messages))
        from apps.orders.services import payment_split, worker_service_percentage
        net, worker_amount, manager_amount = payment_split(order)
        preview_values = None
        if preview:
            preview_values = {
                **preview,
                "amount": Decimal(preview["amount"]),
                "expenses": Decimal(preview["expenses"]),
                "net": Decimal(preview["net"]),
                "worker_amount": Decimal(preview["worker_amount"]) if preview["worker_amount"] is not None else None,
                "company_amount": Decimal(preview["company_amount"]) if preview["company_amount"] is not None else None,
                "worker_percentage": Decimal(preview["worker_percentage"]) if preview["worker_percentage"] is not None else None,
            }
        return render(request, "accounts/order.html", {"order": order, "form": form, "payment_form": payment_form, "completion_form": completion_form, "completion_preview": preview_values, "payment_net": net, "payment_worker": worker_amount, "payment_manager": manager_amount, "payment_percentage": worker_service_percentage(order.employee, order.service)})

@roles_required("manager", "operator")
def edit_record(request, kind, pk=None):
    forms = {"client": ClientForm, "lead": LeadForm, "service": ServiceForm}
    if kind == "service" and not allowed(request.user, "manager"):
        raise PermissionDenied
    if kind not in forms:
        raise PermissionDenied
    form_class = forms[kind]
    obj = get_object_or_404(form_class._meta.model, pk=pk) if pk else None
    if request.method == "POST" and request.POST.get("action") == "lost":
        if kind != "lead" or obj is None:
            raise PermissionDenied
        reason = request.POST.get("reason", "").strip()
        with transaction.atomic():
            lead = Lead.objects.select_for_update().get(pk=obj.pk)
            if Order.objects.filter(lead=lead).exists() or lead.status == "converted":
                messages.error(request, "Уже создан заказ. Закрыть его как неудачную сделку нельзя.")
            elif lead.status == "lost":
                messages.info(request, "Сделка уже закрыта.")
            elif not reason or len(reason) > 300:
                messages.error(request, "Укажите причину отказа (до 300 символов).")
            else:
                from django.utils import timezone
                lead.status = "lost"
                lead.lost_reason = reason
                lead.lost_at = timezone.now()
                lead.lost_by = request.user
                lead.save(update_fields=("status", "lost_reason", "lost_at", "lost_by"))
                TelegramNotice.objects.filter(lead=lead, active=True).update(active=False)
                messages.success(request, "Обращение закрыто как неудачная сделка.")
        return redirect("kanban")
    if request.method == "POST" and request.POST.get("action") == "convert":
        if kind != "lead" or obj is None:
            raise PermissionDenied
        try:
            order, _ = convert_lead(obj.pk, actor=request.user)
            if request.POST.get("return_to") == "kanban":
                return redirect("kanban")
            return redirect("order-detail", pk=order.pk)
        except ValidationError as error:
            messages.error(request, "; ".join(error.messages))
        form = form_class(instance=obj)
    else:
        form = form_class(request.POST or None, instance=obj)
        form.instance._history_actor = request.user
        if request.method == "POST" and form.is_valid():
            form.save()
            return redirect("operator-home")
    return render(request, "accounts/form.html", {"form": form, "kind": kind, "record": obj})

@roles_required("manager")
def employees(request, pk=None):
    obj = get_object_or_404(User, pk=pk, is_superuser=False) if pk else None
    telegram_link = None
    if request.method == "POST" and request.POST.get("action") == "telegram_link":
        if obj is None or obj.role != "worker" or not obj.is_active:
            raise PermissionDenied
        import os, uuid
        from datetime import timedelta
        from django.utils import timezone
        from apps.leads.telegram import TelegramAPI, TelegramError
        try:
            username = TelegramAPI(os.environ.get("TELEGRAM_BOT_TOKEN", "")).call("getMe")["username"]
            with transaction.atomic():
                obj = User.objects.select_for_update().get(pk=obj.pk)
                if obj.telegram_chat_id:
                    messages.error(request, "Telegram уже подключён. Для смены аккаунта сначала очистите Telegram ID и сохраните сотрудника.")
                else:
                    obj.telegram_link_token = uuid.uuid4()
                    obj.telegram_link_expires = timezone.now() + timedelta(hours=24)
                    obj.save(update_fields=("telegram_link_token", "telegram_link_expires"))
                    telegram_link = f"https://t.me/{username}?start=link_{obj.telegram_link_token.hex}"
        except TelegramError as error:
            messages.error(request, str(error))
        form = EmployeeForm(instance=obj)
    else:
        form = EmployeeForm(request.POST or None, instance=obj)
    if request.method == "POST" and form.is_valid():
        # Avoid locking the current manager out through their own role editor.
        if obj and obj.pk == request.user.pk and (form.cleaned_data["role"] != "manager" or not form.cleaned_data["is_active"]):
            form.add_error("role", "Нельзя отключить или изменить собственную роль.")
        else:
            saved = form.save()
            if saved.role == User.Role.WORKER:
                saved.service_cities.set(City.objects.filter(pk__in=request.POST.getlist("service_cities"), is_active=True))
                from decimal import Decimal, InvalidOperation
                for service in Service.objects.all():
                    raw_rate = request.POST.get(f"rate_{service.pk}", "").strip()
                    active = request.POST.get(f"active_{service.pk}") == "on"
                    if not raw_rate:
                        WorkerServiceRate.objects.filter(worker=saved, service=service).delete()
                        continue
                    try:
                        rate = Decimal(raw_rate)
                    except InvalidOperation:
                        form.add_error(None, f"Ставка для услуги «{service.name}» должна быть числом.")
                        break
                    if rate < 0 or rate > 100:
                        form.add_error(None, f"Ставка для услуги «{service.name}» должна быть от 0 до 100%.")
                        break
                    WorkerServiceRate.objects.update_or_create(
                        worker=saved, service=service,
                        defaults={"worker_percentage": rate, "active": active},
                    )
                else:
                    return redirect("employees")
            else:
                WorkerServiceRate.objects.filter(worker=saved).delete()
                return redirect("employees")
    rate_rows = []
    if obj and obj.role == User.Role.WORKER:
        rates = {rate.service_id: rate for rate in WorkerServiceRate.objects.filter(worker=obj)}
        rate_rows = [{"service": service, "rate": rates.get(service.pk)} for service in Service.objects.all()]
    return render(request, "accounts/employees.html", {
        "form": form, "employee_record": obj, "telegram_link": telegram_link,
        "employees": User.objects.filter(is_superuser=False).order_by("username"),
        "rate_rows": rate_rows,
        "cities": City.objects.filter(is_active=True),
    })


@roles_required("manager", "operator")
def resolve_call(request, pk):
    from apps.leads.models import IncomingCall
    from .forms import CallResolveForm
    with transaction.atomic():
        call = get_object_or_404(IncomingCall.objects.select_for_update(), pk=pk)
        if call.lead_id:
            return redirect("record-edit", kind="lead", pk=call.lead_id)
        form = CallResolveForm(request.POST or None, phone=call.phone)
        if request.method == "POST" and form.is_valid():
            call.client = form.cleaned_data["client"]
            call.lead = Lead.objects.create(title="Входящий звонок", client=call.client, source="Телефон")
            call.save(update_fields=("client", "lead"))
            return redirect("record-edit", kind="lead", pk=call.lead_id)
        return render(request, "accounts/form.html", {"form": form, "kind": "client"})


@roles_required("worker")
def assigned_lead_detail(request, pk):
    lead = get_object_or_404(Lead.objects.select_related("client", "service").filter(
        employee=request.user, order__isnull=True, status__in=("new", "assigned", "in_progress"),
    ), pk=pk)
    return render(request, "accounts/assigned_lead.html", {"lead": lead})


@roles_required("worker")
def worker_history(request):
    from django.core.paginator import Paginator
    orders = Order.objects.filter(Q(status__in=("paid", "cancelled")) | (Q(status="completed") & (Q(is_free=True) | Q(repeat_of__isnull=False))), employee=request.user).select_related("client", "service").order_by("-paid_at", "-pk")
    return render(request, "accounts/worker_history.html", {"page_obj": Paginator(orders, 20).get_page(request.GET.get("page"))})


@roles_required("manager", "operator", "worker", "curator")
def profile(request):
    if request.method != "GET":
        from django.http import HttpResponseNotAllowed
        return HttpResponseNotAllowed(["GET"])
    from decimal import Decimal
    from apps.orders.reporting import order_breakdown, financial_summary, local_day_bounds, worker_dashboard
    from django.utils import timezone
    role = role_of(request.user)
    orders = Order.objects.none()
    dashboard_orders = Order.objects.none()
    workers = User.objects.filter(role="worker").order_by("first_name", "last_name", "username")
    selected_worker = None
    if role == "worker":
        orders = Order.objects.filter(employee=request.user)
        dashboard_orders = orders
    elif role == "curator":
        orders = Order.objects.filter(employee__curator=request.user)
        dashboard_orders = orders
    elif role == "manager":
        dashboard_orders = Order.objects.filter(employee__role="worker")
        selected_worker_id = request.GET.get("worker")
        if selected_worker_id and selected_worker_id.isdigit():
            selected_worker = workers.filter(pk=int(selected_worker_id)).first()
        orders = dashboard_orders.filter(employee=selected_worker) if selected_worker else dashboard_orders
    orders = orders.select_related("client", "service", "employee").order_by("-created_at")
    dashboard_orders = dashboard_orders.select_related("client", "service", "employee").order_by("-created_at")
    paid = dashboard_orders.filter(status="paid", repeat_of__isnull=True)
    total = paid.aggregate(total=Sum("amount"))["total"] or Decimal("0")
    earnings = Decimal("0")
    expenses_total = paid.aggregate(total=Sum("expenses"))["total"] or Decimal("0")
    if role in ("worker", "manager"):
        from apps.orders.services import payment_split
        earnings = Decimal("0")
        for order in paid.select_related("employee"):
            net, worker_amount, manager_amount = payment_split(order)
            if worker_amount is None:
                earnings = None
                break
            earnings += worker_amount if role == "worker" else manager_amount
    day_start, day_end = local_day_bounds(timezone.localdate())
    today_orders = dashboard_orders.filter(
        status="paid", repeat_of__isnull=True,
        paid_at__gte=day_start, paid_at__lt=day_end,
    )
    if role == "worker":
        dashboard = worker_dashboard(request.user)
        dashboard["today"] = financial_summary(today_orders)
    elif role == "manager":
        dashboard = {"today": financial_summary(today_orders), "date": timezone.localdate()}
    else:
        dashboard = None
    if dashboard:
        dashboard["date"] = timezone.localdate()
    history = request.GET.get("history", "")
    history_title = {"today": "История за сегодня", "earnings": "История начислений"}.get(history, "Мои заказы")
    if history in ("today", "earnings"):
        orders = today_orders if history == "today" else paid
        if selected_worker:
            orders = orders.filter(employee=selected_worker)
    from django.core.paginator import Paginator
    page_obj = Paginator(orders, 20).get_page(request.GET.get("page"))
    orders = list(page_obj.object_list)
    for order in orders:
        if order.status == "paid":
            order.financial_breakdown = order_breakdown(order)
    return render(request, "accounts/profile.html", {
        "orders": orders,
        "workers": workers if role == "manager" else None,
        "selected_worker": selected_worker,
        "history": history, "history_title": history_title, "page_obj": page_obj,
        "show_earnings": role in ("worker", "curator", "manager"), "paid_total": total,
        "earnings": earnings, "expenses_total": expenses_total, "net_total": total - expenses_total,
        "dashboard": dashboard,
    })


@roles_required("manager", "worker")
def settlement_shift_detail(request, pk):
    from apps.orders.models import SettlementShift
    from apps.orders.reporting import closed_shift_detail
    queryset = SettlementShift.objects.select_related("worker", "closed_by")
    if role_of(request.user) == "worker":
        queryset = queryset.filter(worker=request.user)
    shift = get_object_or_404(queryset, pk=pk)
    return render(request, "accounts/settlement_shift_detail.html", closed_shift_detail(shift))


@roles_required("manager", "worker")
def settlement_export(request, pk=None):
    import openpyxl
    from decimal import Decimal
    from django.http import HttpResponse, HttpResponseBadRequest
    from django.utils import timezone
    from apps.orders.reporting import export_rows, parse_export_date

    is_manager = allowed(request.user, "manager")
    if is_manager:
        worker_id = str(pk or request.GET.get("worker", ""))
        if len(worker_id) > 19 or not worker_id.isascii() or not worker_id.isdigit() or not 0 < int(worker_id) <= 9223372036854775807:
            return HttpResponseBadRequest("Выберите мастера.")
        worker = get_object_or_404(User, pk=int(worker_id), role="worker")
    else:
        if pk and pk != request.user.pk:
            raise PermissionDenied
        worker = request.user
    date_from = parse_export_date(request.GET.get("date_from"))
    date_to = parse_export_date(request.GET.get("date_to"))
    if request.GET.get("date_from") and date_from is None or request.GET.get("date_to") and date_to is None:
        return HttpResponseBadRequest("Дата должна быть указана в формате ГГГГ-ММ-ДД.")
    if date_from and date_to and date_from > date_to:
        return HttpResponseBadRequest("Дата начала не может быть позже даты окончания.")
    try:
        rows = export_rows(worker, date_from, date_to)
    except ValidationError as error:
        return HttpResponseBadRequest("; ".join(error.messages))
    except OverflowError:
        return HttpResponseBadRequest("Выберите дату окончания раньше 9999-12-31.")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Финансы"
    headers = ["Дата", "Номер заказа", "Тип операции", "Стоимость услуг", "Расходы", "После расходов", "Процент мастера", "Доля мастера", "Доля компании", "Перевод", "Остаток", "Комментарий"]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            timezone.localtime(row["date"]).strftime("%Y-%m-%d %H:%M") if row["date"] else "",
            row["order_number"], row["operation"], row["service_amount"], row["expenses"],
            row["net_amount"], row["worker_percentage"], row["worker_amount"],
            row["company_amount"], row["transfer_amount"], row["remaining_balance"], row["comment"],
        ])
        # Preserve user comments as text even when they begin with '='.
        sheet.cell(sheet.max_row, 12).data_type = "s"
    totals_row = len(rows) + 3
    sheet.cell(totals_row, 1, "Итого по операциям")
    for column, key in ((4, "service_amount"), (5, "expenses"), (6, "net_amount"), (8, "worker_amount"), (9, "company_amount"), (10, "transfer_amount")):
        sheet.cell(totals_row, column, sum((row[key] or Decimal("0") for row in rows), Decimal("0")))
    opening = rows[0]["remaining_balance"] if rows and rows[0]["operation"] == "Входящий остаток" else Decimal("0")
    closing = rows[-1]["remaining_balance"] if rows else opening
    summary_rows = [
        ("Входящий остаток", opening),
        ("Переводы получены", sum((row["transfer_amount"] for row in rows), Decimal("0"))),
        ("Остаток на конец периода", closing),
    ]
    for offset, (label, value) in enumerate(summary_rows, start=1):
        sheet.cell(totals_row + offset, 1, label)
        sheet.cell(totals_row + offset, 2, value)
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 18
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="worker-{worker.pk}-financial-history.xlsx"'
    workbook.save(response)
    return response


@roles_required("manager", "worker")
def settlements(request, pk=None):
    import uuid
    from django import forms
    from apps.orders.settlements import totals, settle
    is_manager = allowed(request.user, "manager")
    if not is_manager:
        if pk is not None and pk != request.user.pk:
            raise PermissionDenied
        pk = request.user.pk
    workers = User.objects.filter(role="worker").order_by("username")
    worker = get_object_or_404(workers, pk=pk) if pk else None
    recent_transfers = worker.settlement_transfers.order_by("-created_at", "-pk")[:5] if worker else []
    recent_shifts = worker.closed_shifts.order_by("-closed_at", "-pk")[:5] if worker else []
    context = {
        "worker": worker,
        "workers": workers,
        "recent_transfers": recent_transfers,
        "recent_shifts": recent_shifts,
        "form": None,
        "can_manage": is_manager,
    }
    class TransferForm(forms.Form):
        amount = forms.DecimalField(label="Получено от мастера (KZT)", max_digits=14, decimal_places=2, min_value=0.01)
        comment = forms.CharField(label="Комментарий", max_length=300, required=False)
        request_id = forms.UUIDField(widget=forms.HiddenInput)
    form = TransferForm(request.POST if request.method == "POST" and request.POST.get("action") == "transfer" else None,
                        initial={"request_id": uuid.uuid4()})
    if request.method == "POST":
        if not is_manager or worker is None:
            raise PermissionDenied
        action = request.POST.get("action")
        try:
            if action == "close":
                settle(worker.pk, request.user, action)
            elif action == "transfer" and form.is_valid():
                settle(worker.pk, request.user, action, **form.cleaned_data)
            else:
                context.update({"summary": totals(worker), "form": form})
                return render(request, "accounts/settlements.html", context)
            return redirect("worker-settlement", pk=worker.pk)
        except ValidationError as error:
            messages.error(request, "; ".join(error.messages))
    context.update({"summary": totals(worker) if worker else None, "form": form})
    return render(request, "accounts/settlements.html", context)
